from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Aditya\Downloads\ui-dbms")
TEMPLATE = Path(r"C:\Users\Aditya\Downloads\04. Project Report.docx")
OUTPUT = ROOT / "Course_Feedback_Project_Report.docx"
TMP = ROOT / "tmp" / "docx_report"
ER_IMAGE = TMP / "course_feedback_er_diagram.png"


def clear_document(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def set_font(run, name="Times New Roman", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", align=None, size=12, bold=False, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold)
    return p


def heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=14 if level == 1 else 12, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_font(r, size=11)
    return p


def code_block(doc, text):
    for line in text.strip().splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        set_font(r, name="Courier New", size=9)


def set_cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.autofit = True
    hdr = tbl.rows[0].cells
    for i, text in enumerate(headers):
        set_cell_text(hdr[i], text, bold=True, color="FFFFFF", size=10)
        shade_cell(hdr[i], "1F4E79")
        set_cell_margins(hdr[i])
        if widths:
            hdr[i].width = Inches(widths[i])

    for row_values in rows:
        cells = tbl.add_row().cells
        for i, text in enumerate(row_values):
            set_cell_text(cells[i], text, size=9.5)
            set_cell_margins(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
    paragraph(doc, "", space_after=2)
    return tbl


def create_er_diagram():
    TMP.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1600, 950), "white")
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("arial.ttf", 26)
        small = ImageFont.truetype("arial.ttf", 20)
        title = ImageFont.truetype("arialbd.ttf", 30)
    except OSError:
        font = ImageFont.load_default()
        small = ImageFont.load_default()
        title = ImageFont.load_default()

    boxes = {
        "Admin": (70, 80, 310, 190),
        "Department": (430, 80, 720, 210),
        "Faculty": (860, 80, 1130, 210),
        "Course": (1220, 80, 1510, 210),
        "Student": (70, 380, 350, 520),
        "Teaching Assignment": (500, 360, 860, 540),
        "Enrollment": (990, 380, 1260, 520),
        "Evaluation": (1300, 380, 1540, 520),
        "Evaluation Question": (500, 700, 860, 860),
        "Feedback": (1030, 700, 1300, 860),
    }
    attrs = {
        "Admin": ["AdminID PK", "Email UNIQUE"],
        "Department": ["DepartmentID PK", "AdminID FK", "HOD FK"],
        "Faculty": ["FacultyID PK", "DepartmentID FK"],
        "Course": ["CourseID PK", "CourseCode UNIQUE"],
        "Student": ["StudentID PK", "DepartmentID FK"],
        "Teaching Assignment": ["AssignmentID PK", "CourseID FK", "FacultyID FK", "Section"],
        "Enrollment": ["EnrollmentID PK", "StudentID FK", "AssignmentID FK"],
        "Evaluation": ["EvaluationID PK", "EnrollmentID UNIQUE FK", "Comment"],
        "Evaluation Question": ["QuestionID PK", "MaxScore", "IsActive"],
        "Feedback": ["EvaluationID PK/FK", "QuestionID PK/FK", "Score"],
    }
    lines = [
        ("Admin", "Department", "1 manages M"),
        ("Department", "Faculty", "1 has M"),
        ("Department", "Course", "1 offers M"),
        ("Department", "Student", "1 has M"),
        ("Faculty", "Teaching Assignment", "1 teaches M"),
        ("Course", "Teaching Assignment", "1 appears in M"),
        ("Teaching Assignment", "Enrollment", "1 has M"),
        ("Student", "Enrollment", "1 has M"),
        ("Enrollment", "Evaluation", "1 has 0..1"),
        ("Evaluation", "Feedback", "1 has M"),
        ("Evaluation Question", "Feedback", "1 used in M"),
    ]

    draw.text((450, 18), "Course Feedback and Faculty Evaluation System ER Diagram", fill="black", font=title)

    def center(name):
        x1, y1, x2, y2 = boxes[name]
        return ((x1 + x2) // 2, (y1 + y2) // 2)

    for a, b, label in lines:
        ax, ay = center(a)
        bx, by = center(b)
        draw.line((ax, ay, bx, by), fill="#6B7280", width=3)
        lx, ly = (ax + bx) // 2, (ay + by) // 2
        draw.rectangle((lx - 78, ly - 18, lx + 78, ly + 18), fill="white")
        draw.text((lx - 70, ly - 12), label, fill="#1F2937", font=small)

    for name, (x1, y1, x2, y2) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=12, outline="#1F4E79", width=4, fill="#F3F8FC")
        draw.rectangle((x1, y1, x2, y1 + 42), fill="#1F4E79")
        draw.text((x1 + 14, y1 + 8), name, fill="white", font=font)
        y = y1 + 52
        for attr in attrs[name]:
            draw.text((x1 + 14, y), attr, fill="#111827", font=small)
            y += 26

    img.save(ER_IMAGE)


def add_cover(doc):
    paragraph(doc, "23CSE202 DATABASE MANAGEMENT SYSTEMS", WD_ALIGN_PARAGRAPH.CENTER, 14, True, 3)
    paragraph(doc, "S3 B.Tech CSE", WD_ALIGN_PARAGRAPH.CENTER, 12, True, 3)
    paragraph(doc, "Department of Computer Science and Engineering", WD_ALIGN_PARAGRAPH.CENTER, 12, True, 3)
    paragraph(doc, "Amrita School of Computing, Amritapuri Campus", WD_ALIGN_PARAGRAPH.CENTER, 12, True, 36)
    paragraph(doc, "PROJECT REPORT", WD_ALIGN_PARAGRAPH.CENTER, 18, True, 36)
    paragraph(doc, "Group No: A9", size=12, bold=True)
    paragraph(doc, "Team Members:", size=12, bold=True)
    table(
        doc,
        ["Roll No", "Student Name"],
        [
            ["AM.SC.U4CSE25018", "Aditya Sahu"],
            ["AM.SC.U4CSE25058", "V Chetan Reddy"],
            ["AM.SC.U4CSE25074", "Syed Zaheer Ahmed"],
            ["AM.SC.U4CSE25041", "M Bhargav"],
        ],
        [2.1, 4.2],
    )
    paragraph(doc, "Project Title: Course Feedback and Faculty Evaluation System", size=12, bold=True)
    doc.add_page_break()


def add_contents(doc):
    heading(doc, "CONTENTS", 1)
    table(
        doc,
        ["Sl No", "Chapter", "Page No"],
        [
            ["1", "Abstract", ""],
            ["2", "Introduction", ""],
            ["3", "Requirements", ""],
            ["4", "UI Design", ""],
            ["5", "Database Design", ""],
            ["6", "Database Selection and Justification", ""],
            ["7", "Implementation", ""],
            ["8", "Testing and Validation", ""],
            ["9", "Results and Screenshots", ""],
            ["10", "Challenges Faced and Solutions", ""],
            ["11", "Conclusion and Future Enhancements", ""],
            ["12", "References", ""],
        ],
        [0.8, 4.6, 1.0],
    )
    doc.add_page_break()


def add_report(doc):
    heading(doc, "1 Abstract", 1)
    paragraph(doc, "The Course Feedback and Faculty Evaluation System addresses the difficulty of collecting, storing, and analysing student feedback for courses and faculty members using manual forms or disconnected online forms. The objective of the project is to provide a secure database-backed application where students submit feedback only for enrolled courses, administrators manage academic master data, and faculty or administrators view meaningful aggregate reports. The proposed solution uses PostgreSQL for normalized relational storage, Node.js and Express for the backend API, and a plain HTML, CSS, and JavaScript frontend for a lightweight user interface. Major database features include primary and foreign key constraints, unique constraints to prevent duplicate enrolment and duplicate evaluation, check constraints for valid IDs and scores, password hashing using pgcrypto, deferred constraints for the department HOD relationship, triggers for score validation, and report queries for course, faculty, department, and student feedback summaries. The expected outcome is a reliable feedback system that improves data integrity, reduces repeated manual work, protects student confidentiality in reports, and supports accurate decision making for academic improvement.")

    heading(doc, "2 Introduction", 1)
    paragraph(doc, "Course feedback is an important academic process because it helps departments understand teaching quality, course organization, laboratory effectiveness, student satisfaction, and areas that need improvement. In many colleges, feedback is collected through paper forms or simple online forms. These approaches are easy to start but difficult to maintain because staff must manually check duplicate responses, consolidate scores, and prepare reports. They also make it difficult to ensure that only enrolled students submit feedback for a specific course offering.")
    paragraph(doc, "The proposed system focuses on the course feedback workflow used in a B.Tech CSE academic environment. Students log in using their institutional student ID and name, select one of their enrolled courses, answer a fixed set of evaluation questions on a 1 to 5 scale, and submit an optional comment. Administrators log in with an email and password, maintain student, faculty, and course data, and view dashboards and reports. Faculty members are represented in the database design so that individual faculty reports can be generated without exposing student identities.")
    paragraph(doc, "The existing scenario has several limitations: feedback can be duplicated, unauthorised students can rate courses they did not study, reports are time consuming to prepare, and historical course offerings can be mixed up when the same course is taught by different faculty members in different semesters. The proposed database application addresses these issues by separating courses from teaching assignments, linking students through enrolment, storing one evaluation per enrolment, and storing question-wise feedback responses in a normalized structure.")

    heading(doc, "3 Requirements", 1)
    heading(doc, "3.1 Functional Requirements", 2)
    table(doc, ["Requirement", "Description"], [
        ["Authentication", "Students log in with Student ID and name; admins log in with email and password hash verification."],
        ["Student Feedback Submission", "Students select an enrolled course, answer active evaluation questions, and submit an optional comment."],
        ["Master Data Management", "Admins add and update student, faculty, and course records."],
        ["Teaching Assignment Setup", "The database stores course, faculty, academic year, term, semester, and section for each actual course offering."],
        ["Enrolment Control", "Only students enrolled in a teaching assignment can submit feedback for that course offering."],
        ["Duplicate Prevention", "A student can enrol only once in the same assignment and can submit only one evaluation for one enrolment."],
        ["Report Generation", "The system supports course averages, question-wise averages, faculty reports, department summaries, recent feedback, and pending feedback lists."],
        ["Seed and Reset Support", "Demo data can be loaded using seed.sql and reset through a backend API endpoint."],
    ], [2.2, 4.2])
    heading(doc, "3.2 Non Functional Requirements", 2)
    table(doc, ["Quality", "Requirement"], [
        ["Performance", "Indexes on foreign keys support joins used by dashboards and reports."],
        ["Security", "Passwords are stored as hashes using PostgreSQL pgcrypto; reports for faculty avoid student identity."],
        ["Reliability", "Transactions are used for feedback submission so evaluation and response rows are saved atomically."],
        ["Usability", "The frontend is plain HTML, CSS, and JavaScript with simple role-based navigation."],
        ["Scalability", "The schema separates master data, assignments, enrolments, evaluations, and responses so new terms and questions can be added without redesign."],
        ["Maintainability", "UI logic is separated from API/data logic, and SQL scripts are separated into schema, seed, reports, and tests."],
    ], [1.7, 4.7])

    heading(doc, "4 UI Design", 1)
    paragraph(doc, "The user interface is named Feedback Hub and is implemented with vanilla HTML, CSS, and JavaScript. The same frontend can run as a local prototype or connect to the Express API. The interface has two main roles: student and admin.")
    table(doc, ["Screen", "Purpose", "Main User Interaction"], [
        ["Login Screen", "Allows role selection between Student and Admin.", "Students enter Student ID and full name; admins enter email and password."],
        ["Student Feedback Form", "Collects course feedback.", "Student selects a course, answers evaluation sliders, writes an optional comment, and submits."],
        ["My Feedbacks", "Shows feedback submitted by the logged-in student.", "Student reviews submitted course, average score, date, and comment."],
        ["Profile", "Displays student details.", "Student views ID, department, semester, and feedback count."],
        ["Admin Dashboard", "Shows summary metrics and course rating chart.", "Admin scans totals, average rating, recent feedback, and course performance."],
        ["Students, Faculty, Courses", "Provides add and update forms.", "Admin selects rows, loads existing details, edits fields, and saves changes."],
        ["Feedback and Reports", "Shows feedback summaries and course reports.", "Admin refreshes results, exports feedback, and generates question-wise course reports."],
    ], [1.5, 2.1, 2.8])
    paragraph(doc, "Navigation uses a sidebar in both portals. Student users move between Submit Feedback, My Feedbacks, and Profile. Admin users move between Dashboard, Students, Faculty, Courses, Feedback, and Reports. The frontend calls the backend API when available and falls back to local demo data when the server is offline.")

    heading(doc, "5 Database Design", 1)
    heading(doc, "5.1 ER Diagram", 2)
    paragraph(doc, "The ER diagram below shows the main entities and relationships. Department is the organizational root; teaching assignment represents an actual course offering; enrolment connects students to offerings; evaluation and feedback store one submitted form and its question-wise responses.")
    create_er_diagram()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ER_IMAGE), width=Inches(6.5))
    paragraph(doc, "Figure 1. ER diagram for the Course Feedback and Faculty Evaluation System.", WD_ALIGN_PARAGRAPH.CENTER, 10, False)
    heading(doc, "5.2 Entities and Attributes", 2)
    table(doc, ["Entity", "Important Attributes"], [
        ["Admin", "AdminID, name fields, email, hashed_password, created_at"],
        ["Department", "DepartmentID, DepartmentName, AdminID, HODFacultyID"],
        ["Student", "StudentID, name fields, email, hashed_password, semester, DepartmentID"],
        ["Faculty", "FacultyID, name fields, email, hashed_password, designation, DepartmentID"],
        ["Course", "CourseID, CourseCode, CourseName, SubjectLabel, Credits, DepartmentID"],
        ["Teaching Assignment", "AssignmentID, CourseID, FacultyID, semester, academic_year, term, section"],
        ["Enrollment", "EnrollmentID, StudentID, AssignmentID, EnrollmentDate"],
        ["Evaluation Question", "QuestionID, QuestionText, MaxScore, IsActive, DisplayOrder"],
        ["Evaluation", "EvaluationID, EnrollmentID, SubmissionDate, Status, Comment"],
        ["Feedback", "EvaluationID, QuestionID, Score, Remark"],
    ], [1.7, 4.7])
    heading(doc, "5.3 Relationships and Cardinalities", 2)
    table(doc, ["Relationship", "Cardinality", "Implementation"], [
        ["Admin manages Department", "One admin can manage many departments.", "department.admin_id references admin.admin_id."],
        ["Department has Students, Faculty, Courses", "One department can have many child records.", "Foreign keys in student, faculty, and course."],
        ["Department has HOD", "Each department has one HOD; one faculty member heads at most one department.", "department.hod_faculty_id is unique and uses a deferred same-department foreign key."],
        ["Course and Faculty form Teaching Assignment", "One course and one faculty member create one offering.", "teaching_assignment stores course, faculty, year, term, semester, and section."],
        ["Student enrolls in Teaching Assignment", "Many students can join many assignments through enrollment.", "enrollment resolves the many-to-many relationship."],
        ["Enrollment has Evaluation", "At most one evaluation exists for each enrollment.", "evaluation.enrollment_id is unique."],
        ["Evaluation answers Questions", "One evaluation contains one response per question.", "feedback has composite primary key on evaluation_id and question_id."],
    ], [2.0, 1.6, 2.8])
    heading(doc, "5.4 Relational Schema", 2)
    code_block(doc, """
Admin(AdminID, FirstName, MiddleName, LastName, Email, HashedPassword, CreatedAt)
Admin_Phone(AdminID, Phone)
Department(DepartmentID, DepartmentName, AdminID, HODFacultyID)
Student(StudentID, FirstName, MiddleName, LastName, Email, HashedPassword, Semester, DepartmentID)
Student_Phone(StudentID, Phone)
Faculty(FacultyID, FirstName, MiddleName, LastName, Email, HashedPassword, Designation, DepartmentID)
Faculty_Phone(FacultyID, Phone)
Course(CourseID, CourseCode, CourseName, SubjectLabel, Credits, DepartmentID)
Teaching_Assignment(AssignmentID, CourseID, FacultyID, Semester, AcademicYear, Term, Section)
Enrollment(EnrollmentID, StudentID, AssignmentID, EnrollmentDate)
Evaluation_Question(QuestionID, QuestionText, MaxScore, IsActive, DisplayOrder)
Evaluation(EvaluationID, EnrollmentID, SubmissionDate, Status, Comment)
Feedback(EvaluationID, QuestionID, Score, Remark)
""")
    heading(doc, "5.5 Constraints", 2)
    paragraph(doc, "The schema includes primary keys for all main tables, foreign keys for referential integrity, unique constraints on candidate keys such as email and course code, check constraints for semester, score, course code format, student ID format, and trigger validation to ensure feedback scores do not exceed the selected question's maximum score. Foreign-key indexes are created for efficient joins and delete checks.")
    heading(doc, "5.6 Functional Dependencies and Normalization", 2)
    table(doc, ["Relation", "Main Functional Dependencies"], [
        ["Admin", "AdminID -> name, email, hashed_password; Email -> AdminID"],
        ["Department", "DepartmentID -> DepartmentName, AdminID, HODFacultyID; HODFacultyID -> DepartmentID"],
        ["Student", "StudentID -> name, email, semester, DepartmentID; Email -> StudentID"],
        ["Faculty", "FacultyID -> name, email, designation, DepartmentID; Email -> FacultyID"],
        ["Course", "CourseID -> CourseCode, CourseName, Credits, DepartmentID; CourseCode -> CourseID"],
        ["Teaching Assignment", "AssignmentID -> CourseID, FacultyID, AcademicYear, Term, Section"],
        ["Enrollment", "EnrollmentID -> StudentID, AssignmentID, EnrollmentDate; StudentID + AssignmentID -> EnrollmentID"],
        ["Evaluation", "EvaluationID -> EnrollmentID, SubmissionDate, Status, Comment; EnrollmentID -> EvaluationID"],
        ["Feedback", "EvaluationID + QuestionID -> Score, Remark"],
    ], [1.7, 4.7])
    paragraph(doc, "The design satisfies 1NF because all fields are atomic and repeating question scores are separated into the Feedback table. It satisfies 2NF because non-key attributes depend on full primary keys; for example, question text is stored in Evaluation_Question rather than inside Feedback. It satisfies 3NF because department, course, faculty, student, and assignment details are not copied into feedback rows. Candidate keys such as email, course code, enrolment pair, and evaluation enrolment are protected using unique constraints, supporting BCNF-level reasoning for the core relations.")

    heading(doc, "6 Database Selection and Justification", 1)
    paragraph(doc, "PostgreSQL was selected as the DBMS because it provides strong SQL support, reliable transactions, rich constraint handling, triggers, extensions, indexing, and good support for relational database design. The pgcrypto extension supports password hashing inside the database. PostgreSQL also supports deferrable foreign key constraints, which are useful for the circular Department-HOD-Faculty relationship. It is open source, widely documented, and suitable for student projects as well as production applications.")
    table(doc, ["Factor", "Justification"], [
        ["Performance", "Foreign-key indexes and SQL aggregates support dashboard and report queries efficiently."],
        ["Security", "Password hashes can be verified using pgcrypto; row-level reporting can avoid exposing student identities."],
        ["Integrity", "Primary keys, foreign keys, unique constraints, check constraints, and triggers enforce business rules."],
        ["Scalability", "The normalized schema supports new departments, courses, terms, sections, and questions."],
        ["Cost", "PostgreSQL is free and open source."],
        ["Ease of Use", "It works well with pgAdmin, psql, Node.js, and the pg library."],
    ], [1.4, 5.0])

    heading(doc, "7 Implementation", 1)
    paragraph(doc, "The implementation uses PostgreSQL for data storage, Node.js with Express for the backend API, and plain HTML, CSS, and JavaScript for the frontend. The database scripts are separated by purpose: schema.sql creates tables and constraints, seed.sql loads demo data, reports.sql contains report queries, and test_cases.sql validates constraints.")
    heading(doc, "7.1 Project Files", 2)
    table(doc, ["File", "Purpose"], [
        ["schema.sql", "Creates schema, keys, constraints, indexes, triggers, and pgcrypto extension."],
        ["seed.sql", "Loads CSE department, admin, faculty, students, courses, assignments, enrollments, questions, evaluations, and feedback."],
        ["server.js", "Express backend with PostgreSQL connection pool and API routes."],
        ["data.js", "Frontend data adapter that calls API endpoints with local fallback data."],
        ["index.html, style.css, script.js", "Student and admin UI implementation."],
        ["reports.sql", "Reusable SQL report queries."],
        ["test_cases.sql", "PASS/FAIL validation test script."],
    ], [1.7, 4.7])
    heading(doc, "7.2 Database Creation Commands", 2)
    code_block(doc, """
psql -U postgres -d postgres -c "CREATE DATABASE feedback_hub;"
psql -U postgres -d feedback_hub -f schema.sql
psql -U postgres -d feedback_hub -f seed.sql
""")
    heading(doc, "7.3 Important SQL Snippets", 2)
    paragraph(doc, "The following examples show important implementation logic from the schema.")
    code_block(doc, """
ALTER TABLE department
    ADD CONSTRAINT fk_department_hod_same_department
    FOREIGN KEY (department_id, hod_faculty_id)
    REFERENCES faculty(department_id, faculty_id)
    DEFERRABLE INITIALLY DEFERRED;
""")
    code_block(doc, """
CREATE TRIGGER trg_feedback_validate_score
BEFORE INSERT OR UPDATE OF question_id, score ON feedback
FOR EACH ROW
EXECUTE FUNCTION validate_feedback_score();
""")
    heading(doc, "7.4 API Routes", 2)
    table(doc, ["Endpoint", "Method", "Purpose"], [
        ["/api/health", "GET", "Checks database connectivity."],
        ["/api/data", "GET", "Loads admin, questions, students, faculty, courses, and feedback for the UI."],
        ["/api/login", "POST", "Authenticates student or admin."],
        ["/api/feedback", "POST", "Creates evaluation and question-wise feedback rows in one transaction."],
        ["/api/students", "POST", "Adds a new student."],
        ["/api/students/:id", "PUT", "Updates student details."],
        ["/api/teachers", "POST", "Adds a faculty member."],
        ["/api/courses", "POST", "Adds course and teaching assignment."],
        ["/api/reset-seed", "POST", "Reloads demo seed data."],
    ], [1.7, 1.0, 3.7])

    heading(doc, "8 Testing and Validation", 1)
    paragraph(doc, "Testing was planned around both normal functional behaviour and exceptional constraint cases. The test_cases.sql script catches expected failures and prints a PASS or FAIL result table.")
    table(doc, ["Test Case", "Input", "Expected Output", "Actual Output", "Result"], [
        ["Admin login hash", "admin@amrita.edu / admin123", "One matching admin row", "Password hash condition returns one row", "Pass"],
        ["Duplicate student", "Insert existing AM.SC.U4CSE25018", "Primary key violation", "Unique violation raised", "Pass"],
        ["Duplicate enrollment", "Same student and assignment twice", "Unique constraint violation", "Unique violation raised", "Pass"],
        ["Invalid student ID", "BAD-ID", "Check constraint violation", "Check violation raised", "Pass"],
        ["Score above max", "Score 6 for max score 5", "Trigger rejection", "Check violation raised", "Pass"],
        ["Inactive question", "Submit response to inactive question", "Trigger rejection", "Check violation raised", "Pass"],
        ["Duplicate evaluation", "Second evaluation for same enrollment", "Unique constraint violation", "Unique violation raised", "Pass"],
        ["Wrong HOD department", "ECE department with CSE faculty HOD", "Deferred foreign key rejection", "Foreign key violation raised", "Pass"],
    ], [1.35, 1.55, 1.35, 1.35, 0.7])
    paragraph(doc, "The validation confirms that the database enforces entity integrity, referential integrity, domain constraints, duplicate prevention, password hash checking, and important cross-table business rules.")

    heading(doc, "9 Results and Screenshots", 1)
    paragraph(doc, "The completed system supports student feedback submission, admin data management, and report generation using live PostgreSQL data. The seeded database contains 4 students, 3 faculty members, 3 courses, 5 evaluation questions, 4 evaluations, and 20 question-wise feedback rows. The frontend dashboard shows total counts, overall average rating, course-wise averages, and recent feedback comments.")
    table(doc, ["Screenshot", "Caption to Use in Final Submission"], [
        ["Login screen", "Figure 2. Feedback Hub login screen with student and admin role selection."],
        ["Student feedback form", "Figure 3. Student portal showing course selection, rating sliders, and comment box."],
        ["Admin dashboard", "Figure 4. Admin dashboard showing student, faculty, course, and average rating metrics."],
        ["Course management", "Figure 5. Admin course management screen for adding and updating course records."],
        ["pgAdmin tables", "Figure 6. PostgreSQL public schema showing project tables."],
        ["SQL reports", "Figure 7. Output of report query showing course-wise average ratings."],
        ["Test cases", "Figure 8. PASS results from test_cases.sql validating database constraints."],
    ], [2.0, 4.4])
    paragraph(doc, "Screenshots can be inserted at the positions listed above after running the application and SQL scripts on the local machine. Each screenshot should show the complete window area, have a short caption, and demonstrate either a UI workflow or a database result.")

    heading(doc, "10 Challenges Faced and Solutions", 1)
    table(doc, ["Challenge", "Solution"], [
        ["Circular Department-HOD-Faculty relationship", "Created Department first and added a deferrable composite foreign key after Faculty creation."],
        ["Preventing duplicate feedback", "Made evaluation.enrollment_id unique so each enrolment can have at most one submitted evaluation."],
        ["Ensuring only enrolled students submit feedback", "Feedback submission API looks up an enrollment for the selected student and course before insertion."],
        ["Validating score against question maximum", "Added a trigger that checks score against Evaluation_Question.max_score and rejects inactive questions."],
        ["Plain UI to real backend migration", "Kept UI logic separate from data access in data.js and swapped mock calls for API calls."],
        ["Password handling", "Used pgcrypto hashes instead of storing plaintext passwords."],
        ["Report privacy", "Faculty report query uses anonymous aggregate scores and does not select student identity."],
    ], [2.0, 4.4])

    heading(doc, "11 Conclusion and Future Enhancements", 1)
    paragraph(doc, "The project successfully implements a database-backed Course Feedback and Faculty Evaluation System. The schema models departments, students, faculty, courses, teaching assignments, enrollments, evaluation questions, evaluations, and feedback responses in a normalized form. PostgreSQL constraints and triggers enforce important rules such as valid student IDs, valid scores, one evaluation per enrollment, and same-department HOD assignment. The Node.js backend connects the database to a plain HTML, CSS, and JavaScript UI, allowing students to submit feedback and admins to manage data and generate reports.")
    paragraph(doc, "Future enhancements can include JWT-based sessions, separate faculty login screens, role-based authorization middleware, CSV or PDF export of reports, richer analytics dashboards, mobile-responsive improvements, email notifications for pending feedback, cloud deployment, and row-level security for stronger confidentiality. Additional report views can also compare trends across academic years, sections, and departments.")

    heading(doc, "12 References", 1)
    paragraph(doc, "1. PostgreSQL Global Development Group. PostgreSQL Documentation. https://www.postgresql.org/docs/")
    paragraph(doc, "2. Express.js. Express Web Framework Documentation. https://expressjs.com/")
    paragraph(doc, "3. node-postgres. PostgreSQL Client for Node.js Documentation. https://node-postgres.com/")
    paragraph(doc, "4. Amrita School of Computing. 23CSE202 Database Management Systems Project Report Format.")
    paragraph(doc, "5. Course Feedback and Faculty Evaluation System project source files: schema.sql, seed.sql, server.js, data.js, reports.sql, and test_cases.sql.")


def main():
    TMP.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_document(doc)
    section = doc.sections[0]
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    for style_name in ["Normal", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(12)
    add_cover(doc)
    add_contents(doc)
    add_report(doc)
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
